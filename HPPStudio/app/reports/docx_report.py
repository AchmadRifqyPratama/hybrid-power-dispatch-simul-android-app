"""Laporan Studi Kelayakan / Simulasi Dispatch — Word (.docx).

Me-render model blok dari `content.build_report()` sehingga isi & urutan
IDENTIK dengan versi PDF. Memakai gaya Heading bawaan Word (agar Daftar Isi
otomatis berfungsi), header/footer bernomor halaman, sampul berpita brand.
Grafik & persamaan di-embed sebagai gambar (matplotlib).

PENTING: modul ini TIDAK menjalankan/mengubah mesin dispatch — hanya MEMBACA
hasil simulasi melalui `content.py`.
"""
from __future__ import annotations
import io
import re

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from . import common as K
from .common import L
from .content import build_report

CONTENT_CM = 17.0                      # A4 21 cm − margin 2 + 2 cm
FONT = 'Calibri'
_AL = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT,
       'c': WD_ALIGN_PARAGRAPH.CENTER, 'j': WD_ALIGN_PARAGRAPH.JUSTIFY}
TONE = {'good': (K.OK, 'EAF6EE'), 'warn': (K.WARN, 'FDF3E7'),
        'bad': (K.BAD, 'FBEAEA'), 'info': (K.INFO, 'EAF0FB')}
GROUP_BG = 'E3ECE6'


def _rgb(h):
    return RGBColor.from_string(h.lstrip('#'))


# ------------------------------------------------------------- XML util -------
def _shade(cell, hexs):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexs.lstrip('#'))
    tcPr.append(shd)


def _cell_borders(cell, **edges):
    """edges: top/bottom/left/right = (sz, hex) atau None (tanpa garis)."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge, spec in edges.items():
        e = OxmlElement(f'w:{edge}')
        if spec is None:
            e.set(qn('w:val'), 'nil')
        else:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(spec[0]))
            e.set(qn('w:space'), '0'); e.set(qn('w:color'), spec[1].lstrip('#'))
        b.append(e)


def _table_borders(table, inside_h=(4, K.RULE), outer=None):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        spec = inside_h if edge in ('insideH', 'bottom') else outer
        if spec is None:
            e.set(qn('w:val'), 'nil')
        else:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(spec[0]))
            e.set(qn('w:space'), '0'); e.set(qn('w:color'), spec[1].lstrip('#'))
        borders.append(e)
    tblPr.append(borders)


def _cell_margins(table, top=40, bottom=40, left=90, right=90):
    tblPr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for side, v in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa')
        m.append(e)
    tblPr.append(m)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); e.set(qn('w:val'), 'true'); trPr.append(e)


def _widths(table, fracs):
    table.autofit = False
    for row in table.rows:
        for c, f in zip(row.cells, fracs):
            c.width = Cm(CONTENT_CM * f)


def _field(par, code, size=None, color=None):
    run = par.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = code
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = ' '
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    for x in (f1, it, f2, t, f3):
        run._r.append(x)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = _rgb(color)
    return run


def _para_border(p, edge, color, sz=12, space=2):
    pPr = p._p.get_or_add_pPr()
    pb = pPr.find(qn('w:pBdr'))
    if pb is None:
        pb = OxmlElement('w:pBdr'); pPr.append(pb)
    b = OxmlElement(f'w:{edge}')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(sz))
    b.set(qn('w:space'), str(space)); b.set(qn('w:color'), color.lstrip('#'))
    pb.append(b)


# ------------------------------------------------------------- teks -----------
def _runs(par, text, size=None, color=K.INK, bold=False, italic=False):
    """Tambah teks dengan dukungan **tebal**."""
    parts = re.split(r'(\*\*.+?\*\*)', str(text))
    for part in parts:
        if not part:
            continue
        b = part.startswith('**') and part.endswith('**')
        r = par.add_run(part[2:-2] if b else part)
        r.bold = bold or b
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = _rgb(color)
    return par


def _cell_text(cell, text, size=8.5, bold=False, color=K.INK, align='l'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = _AL[align]
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    _runs(p, text, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _setup_styles(doc):
    st = doc.styles
    normal = st['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(K.INK)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color, before, after in (('Heading 1', 16, K.DARK, 6, 10),
                                              ('Heading 2', 12, K.DARK, 12, 4),
                                              ('Heading 3', 12, K.DARK, 12, 4)):
        h = st[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = _rgb(color)
        hr = h.element.get_or_add_rPr()
        f = hr.find(qn('w:rFonts'))
        if f is None:
            f = OxmlElement('w:rFonts'); hr.append(f)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            f.set(qn(a), FONT)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    for name in ('List Bullet', 'List Number'):
        try:
            s = st[name]
            s.font.name = FONT
            s.font.size = Pt(10)
            s.paragraph_format.space_after = Pt(3)
        except KeyError:
            pass


# ------------------------------------------------------------- blok -----------
def _table(doc, b):
    size = (b.get('font') or 8.3) + 0.5
    hdr = b['header']
    fr = b.get('widths') or [1.0 / len(hdr)] * len(hdr)
    al = b['align']
    if b.get('caption'):
        cp = doc.add_paragraph()
        cp.paragraph_format.keep_with_next = True
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(3)
        _runs(cp, b['caption'], size=9, color=K.DARK, bold=True)
    t = doc.add_table(rows=1, cols=len(hdr))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    _table_borders(t)
    _cell_margins(t)
    for j, h in enumerate(hdr):
        c = t.rows[0].cells[j]
        _cell_text(c, h, size - 0.3, True, K.WHITE, al[j])
        _shade(c, K.DARK)
    _repeat_header(t.rows[0])
    tones = b.get('tones') or {}
    groups = b.get('group_rows') or set()
    n = len(b['rows'])
    zebra = 0
    for i, r in enumerate(b['rows']):
        row = t.add_row()
        _no_split(row)
        cells = row.cells
        if i in groups:
            m = cells[0].merge(cells[-1])
            _cell_text(m, r[0], size, True, K.DARK)
            _shade(m, GROUP_BG)
            zebra = 0
            continue
        zebra += 1
        is_total = b.get('total') and i == n - 1
        for j, v in enumerate(r):
            tone = tones.get((i, j))
            col = TONE[tone][0] if tone else K.INK
            _cell_text(cells[j], v, size, bool(is_total or tone), col, al[j])
            if is_total:
                _shade(cells[j], GROUP_BG)
                _cell_borders(cells[j], top=(10, K.DARK))
            elif zebra % 2 == 0:
                _shade(cells[j], K.LIGHT)
    _widths(t, fr)
    if n <= 16:                            # tabel pendek: jangan terpecah
        for row in t.rows[:-1]:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


def _kpis(doc, b):
    items = b['items']
    cols = max(1, min(int(b.get('cols') or 4), len(items)))
    rows = (len(items) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    _table_borders(t, inside_h=None, outer=None)
    _cell_margins(t, 80, 80, 140, 80)
    for k, (lab, val, sub) in enumerate(items):
        c = t.rows[k // cols].cells[k % cols]
        c.text = ''
        _shade(c, K.LIGHT)
        _cell_borders(c, left=(24, K.GREEN), top=(12, 'FFFFFF'), bottom=(12, 'FFFFFF'),
                      right=(12, 'FFFFFF'))
        p1 = c.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        _runs(p1, lab, size=8, color=K.GREY)
        p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
        _runs(p2, val, size=14, color=K.DARK, bold=True)
        if sub:
            p3 = c.add_paragraph(); p3.paragraph_format.space_after = Pt(0)
            _runs(p3, sub, size=7.5, color=K.GREY)
    _widths(t, [1.0 / cols] * cols)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _callout(doc, b):
    fg, bg = TONE.get(b.get('tone'), TONE['info'])
    icon = {'good': '✓', 'warn': '!', 'bad': '✗', 'info': 'i'}.get(b.get('tone'), 'i')
    t = doc.add_table(rows=1, cols=1)
    _table_borders(t, inside_h=None, outer=None)
    _cell_margins(t, 80, 80, 180, 120)
    c = t.rows[0].cells[0]
    _shade(c, bg)
    _cell_borders(c, left=(36, fg))
    c.text = ''
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    _runs(p, f"{icon}  {b['title']}", size=10.5, color=fg, bold=True)
    if b.get('text'):
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        _runs(p2, b['text'], size=9, color=K.INK)
    _widths(t, [1.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _figure(doc, b):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    w = CONTENT_CM * float(b.get('width') or 1.0)
    p.add_run().add_picture(io.BytesIO(b['png']), width=Cm(w))
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(10)
    _runs(cap, b['caption'], size=8.5, color=K.GREY, italic=True)


def _equation(doc, latex):
    buf, w_pt, _h = K.eq_png(latex, fs=11)
    w_in = min(w_pt / 72.0, CONTENT_CM / 2.54 - 0.2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(buf, width=Inches(w_in))


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------- sampul & kaki -----
def _cover(doc, meta):
    band = doc.add_table(rows=1, cols=1)
    _table_borders(band, inside_h=None, outer=None)
    _cell_margins(band, 300, 360, 360, 300)
    c = band.rows[0].cells[0]
    _shade(c, K.DARK)
    _cell_borders(c, bottom=(36, K.GREEN))
    c.text = ''
    p = c.paragraphs[0]
    logo = K.asset('logo_akartha_dark.png') or K.asset('logo_akartha.png')
    if logo:
        try:
            p.add_run().add_picture(logo, width=Cm(4.2))
        except Exception:
            pass
    p.paragraph_format.space_after = Pt(28)
    for text, size, color, bold, after in (
            ('  ·  '.join(a.upper() for a in meta['aspects']), 9.5, '8FD3A8', True, 4),
            (meta['title'], 24, 'FFFFFF', True, 6),
            (meta['subtitle'], 11, 'C9DDD1', False, 10),
            (meta['project'], 14, 'FFFFFF', True, 6)):
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(after)
        _runs(q, text, size=size, color=color, bold=bold)
    _widths(band, [1.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    rows = [(L('Klien', 'Client'), meta['client']), (L('Site', 'Site'), meta['site']),
            (L('Lokasi', 'Location'), meta['location'])]
    if meta.get('coords') and meta['coords'] != '-':
        rows.append((L('Koordinat', 'Coordinates'), meta['coords']))
    rows += [(L('Nomor dokumen', 'Document no.'), meta['doc_no']),
             (L('Revisi', 'Revision'), f"{meta['revision']} — {meta['status']}"),
             (L('Tanggal', 'Date'), meta['date'])]
    t = doc.add_table(rows=0, cols=2)
    _table_borders(t)
    _cell_margins(t, 60, 60, 0, 80)
    for k, v in rows:
        r = t.add_row().cells
        _cell_text(r[0], k, 9, False, K.GREY)
        _cell_text(r[1], v, 9.5, True, K.INK)
    _widths(t, [0.25, 0.75])
    if meta.get('verdict') is not None:
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        _callout(doc, {'tone': 'good' if meta['verdict'] else 'warn', 'title': (
            L('Kesimpulan: layak secara finansial', 'Verdict: financially feasible')
            if meta['verdict'] else
            L('Kesimpulan: marginal — tinjau asumsi', 'Verdict: marginal — review assumptions'))})
    _page_break(doc)


def _header_footer(section, meta):
    section.different_first_page_header_footer = True
    fp = section.first_page_footer.paragraphs[0]
    _runs(fp, f"{meta['prepared']}   ·   {meta['doc_no']}  ·  Rev. {meta['revision']}"
              f"  ·  {meta['footer']}", size=7.5, color=K.GREY)
    hp = section.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM),
                                               alignment=2)       # kanan
    _runs(hp, meta['project'][:70], size=8, color=K.DARK, bold=True)
    hp.add_run('\t')
    _runs(hp, meta['header_right'], size=8, color=K.GREY)
    _para_border(hp, 'bottom', K.GREEN, sz=8, space=4)
    fp2 = section.footer.paragraphs[0]
    fp2.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM), alignment=2)
    _para_border(fp2, 'top', K.RULE, sz=4, space=4)
    _runs(fp2, f"{meta['doc_no']}  ·  Rev. {meta['revision']}  ·  {meta['footer']}",
          size=7.5, color=K.GREY)
    fp2.add_run('\t')
    _runs(fp2, L('Halaman ', 'Page '), size=7.5, color=K.GREY)
    _field(fp2, 'PAGE', 7.5, K.GREY)
    _runs(fp2, ' / ', size=7.5, color=K.GREY)
    _field(fp2, 'NUMPAGES', 7.5, K.GREY)


def _fc(run, kind):
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
    run._r.append(e)


def _toc_field(doc, blocks):
    """Field TOC Word; hasil awal diisi daftar judul (tanpa nomor halaman)
    sehingga tetap terbaca meski field belum di-update. Word akan menawarkan
    pembaruan saat dokumen dibuka (updateFields)."""
    heads = []
    for b in blocks:
        if b['t'] == 'h1':
            num = b['num']
            heads.append((0, f"{num}.  {b['text']}" if num[0].isdigit()
                          else f"{num} — {b['text']}"))
        elif b['t'] == 'h2' and b['num'][0].isdigit():
            heads.append((1, f"{b['num']}  {b['text']}"))
    if not heads:
        heads = [(0, '')]
    for k, (lvl, txt) in enumerate(heads):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(1)
        pf.left_indent = Cm(0.6 * lvl)
        if k == 0:
            r = p.add_run(); _fc(r, 'begin')
            r2 = p.add_run()
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
            it.text = 'TOC \\o "1-2" \\h \\z \\u'
            r2._r.append(it)
            r3 = p.add_run(); _fc(r3, 'separate')
        _runs(p, txt, size=10 if lvl == 0 else 9, color=K.DARK if lvl == 0 else K.INK,
              bold=(lvl == 0))
        if k == len(heads) - 1:
            r4 = p.add_run(); _fc(r4, 'end')


def _update_fields_on_open(doc):
    s = doc.settings.element
    e = OxmlElement('w:updateFields'); e.set(qn('w:val'), 'true')
    s.append(e)


# ================================================================ builder ======
def build_feasibility_docx(path, cfg, result, datetimes=None, multi_result=None,
                           include_technical=True, include_financial=True):
    """Bangun laporan Word di `path` (aspek teknis / finansial sesuai pilihan)."""
    meta, blocks = build_report(cfg, result, datetimes=datetimes,
                                multi_result=multi_result,
                                include_technical=include_technical,
                                include_financial=include_financial)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    _setup_styles(doc)
    _header_footer(sec, meta)
    cp = doc.core_properties
    cp.title = f"{meta['title']} — {meta['project']}"
    cp.author = 'Akartha Energy'
    cp.comments = ''                       # hapus 'generated by python-docx'
    cp.subject = meta['title']
    _cover(doc, meta)
    for b in blocks:
        t = b['t']
        if t == 'pagebreak':
            _page_break(doc)
        elif t == 'doc_control':
            h = doc.add_paragraph()
            _runs(h, L('Kendali Dokumen', 'Document Control'), size=14, color=K.DARK, bold=True)
            _table(doc, {'header': [L('Item', 'Item'), L('Keterangan', 'Details')],
                         'rows': [list(r) for r in b['rows']], 'align': ['l', 'l'],
                         'widths': [0.3, 0.7]})
        elif t == 'toc':
            h = doc.add_paragraph()
            h.paragraph_format.page_break_before = True
            _runs(h, L('Daftar Isi', 'Contents'), size=14, color=K.DARK, bold=True)
            _toc_field(doc, blocks)
            n = doc.add_paragraph()
            _runs(n, L('(Klik kanan → Update Field bila daftar isi belum tampil.)',
                       '(Right-click → Update Field if the table of contents is not shown.)'),
                  size=8, color=K.GREY, italic=True)
        elif t == 'h1':
            num = b['num']
            txt = f"{num}.  {b['text']}" if num[0].isdigit() else f"{num} — {b['text']}"
            h = doc.add_heading(level=1)
            _runs(h, txt, color=K.DARK)
            _para_border(h, 'bottom', K.GREEN, sz=12, space=4)
        elif t == 'h2':
            h = doc.add_heading(level=2 if b['num'][0].isdigit() else 3)
            _runs(h, f"{b['num']}  {b['text']}", color=K.DARK)
        elif t == 'p':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if b.get('keep') or str(b['text']).rstrip().endswith(':'):
                p.paragraph_format.keep_with_next = True
            _runs(p, b['text'], size=10)
        elif t == 'note':
            p = doc.add_paragraph()
            if b.get('keep'):
                p.paragraph_format.keep_with_next = True
            _runs(p, b['text'], size=8.5, color=K.GREY, italic=True)
        elif t in ('bullets', 'numbered'):
            style = 'List Bullet' if t == 'bullets' else 'List Number'
            for it in b['items']:
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:
                    p = doc.add_paragraph()
                _runs(p, it, size=10)
        elif t == 'table':
            _table(doc, b)
        elif t == 'kpis':
            _kpis(doc, b)
        elif t == 'callout':
            _callout(doc, b)
        elif t == 'figure':
            _figure(doc, b)
        elif t == 'eq':
            _equation(doc, b['latex'])
    _update_fields_on_open(doc)
    doc.save(path)
    return path
